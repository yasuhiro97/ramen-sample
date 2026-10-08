#!/usr/bin/env python3
"""輸入事後調査資料の作成ツール。

輸入許可通知書(PDF)から申告番号・輸入許可年月日・仕出人を読み取り、
  1. 管理番号(No.)を採番
  2. 輸入明細一覧(Excel)へ行を追加
  3. 「表紙」(Excel)を作成
してデスクトップの作業フォルダに保存する。

使い方:
  python import_audit.py 許可通知書.pdf [...] --office "林六／東京"            # 海外送金あり
  python import_audit.py 許可通知書.pdf --nosend 着払 --office "林六／東京"     # 海外送金なし(着払/無償/乙仲)
  python import_audit.py フォルダ名 --office ...                                # フォルダ内の PDF すべて

送金日・金額・乙仲業者名など PDF にない項目は、台帳の空欄に手入力する。
"""
import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.styles import PatternFill

HERE = Path(__file__).resolve().parent
TPL_LEDGER = HERE / "templates" / "輸入明細一覧テンプレート.xlsx"
TPL_COVER = HERE / "templates" / "表紙テンプレート.xlsx"
SHEET_SEND = "海外送金あり"
SHEET_NOSEND = "海外送金なし（乙仲・無償・着払）"
HILITE = PatternFill("solid", fgColor="FFFF00")


# ---------- 期・出力先 ----------
def fiscal_period(d: datetime) -> int:
    """4月始まり。2026/4〜2027/3 が第81期。"""
    y = d.year if d.month >= 4 else d.year - 1
    return y - 1945


def period_label(n: int) -> str:
    y = n + 1945
    return f"第{n}期（{y}.4月～{y + 1}.3月）"


def desktop() -> Path:
    home = Path.home()
    for p in (home / "Desktop", home / "OneDrive" / "Desktop", home / "デスクトップ"):
        if p.is_dir():
            return p
    return home


# ---------- PDF 読み取り ----------
_ocr_engine = None
DATE_RE = re.compile(r"(20\d{2})[/年.-](\d{1,2})[/月.-](\d{1,2})")
DECL_RE = re.compile(r"(?<!\d)(\d{3})[ \u3000]?(\d{4})[ \u3000]?(\d{4})(?!\d)")


def ocr_items(pdf: Path):
    """スキャン PDF を OCR し、[(y, x, text)] を上から下の順で返す(1ページ目のみ)。"""
    global _ocr_engine
    try:
        import pypdfium2 as pdfium
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        raise ValueError("スキャンPDFの読み取りには `pip install rapidocr-onnxruntime pypdfium2` が必要です")
    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    page = pdfium.PdfDocument(str(pdf))[0]
    img = page.render(scale=3).to_numpy()
    res, _ = _ocr_engine(img)
    h = img.shape[0]
    return [(b[0][1] / h, b[0][0], t) for b, t, _ in (res or [])]


def norm_name(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", s.upper().replace("LID", "LTD"))


def match_supplier(raw: str, suppliers: list[str]) -> tuple[str, bool]:
    """OCRの崩れた仕出人名を既知の仕入先名に寄せる。(名前, 確実か) を返す。"""
    import difflib
    key = norm_name(raw)
    best, score = raw, 0.0
    for name in suppliers:
        r = difflib.SequenceMatcher(None, key, norm_name(name)).ratio()
        if r > score:
            best, score = name, r
    return (best, True) if score >= 0.85 else (raw, False)


def parse_items(items):
    """items: [(y比率, x, text)]。ラベルが誤認識されても数字・位置で拾う。"""
    items = sorted(items)
    alltext = "\n".join(t for _, _, t in items)
    head = [t for y, _, t in items if y < 0.2]
    decl = next((" ".join(m.groups()) for t in head if (m := DECL_RE.search(t))), None)
    if not decl:
        # 下部の「輸入申告番号等 11桁」
        m = re.search(r"(?<!\d)(\d{3})(\d{4})(\d{4})(?!\d)", alltext)
        decl = " ".join(m.groups()) if m else None
    # 許可日: ページ下部(8割より下)の最初の日付 → なければ上部の申告年月日
    permit = None
    for lo in (0.8, 0.0):
        for y, _, t in items:
            if y >= lo and (m := DATE_RE.search(t)):
                permit = datetime(*map(int, m.groups()))
                break
        if permit:
            break
    # 仕出人: 「出人」を含むラベルの右隣(同じ高さ)の文字列
    shipper = ""
    for y, x, t in items:
        if re.search(r"仕\s*出\s*人|出\s*人$", t) and len(t) <= 6:
            cand = [(xx, tt) for yy, xx, tt in items if abs(yy - y) < 0.012 and xx > x + 50 and re.search(r"[A-Za-z]{3}", tt)]
            if cand:
                shipper = min(cand)[1]
                break
    return decl, permit, shipper


def read_permit(pdf: Path, suppliers: list[str] = ()) -> dict:
    with pdfplumber.open(pdf) as doc:
        text = "\n".join((p.extract_text(layout=True) or "") for p in doc.pages)
    decl = permit = None
    shipper = ""
    if text.strip():  # 文字情報あり
        flat = re.sub(r"[ \u3000]+", " ", text)
        m = re.search(r"申告番号\s*\n?[^\d]*?(\d{3})\s?(\d{4})\s?(\d{4})", flat) or DECL_RE.search(flat)
        decl = " ".join(m.groups()) if m else None
        m = re.search(r"輸入許可日\s*(\d{4})[/年.-](\d{1,2})[/月.-](\d{1,2})", flat)
        permit = datetime(*map(int, m.groups())) if m else None
        m = re.search(r"仕\s*出\s*人\s*[-ー－]?\s*([^\n]+)", text)
        if m:
            shipper = re.sub(r"\s{2,}.*$", "", m.group(1)).strip(" -")
    if not (decl and permit and shipper):  # 足りなければ OCR で補う
        d2, p2, s2 = parse_items(ocr_items(pdf))
        decl, permit, shipper = decl or d2, permit or p2, shipper or s2
    if not decl:
        raise ValueError("申告番号(11桁)が見つかりません。輸入許可通知書ではない可能性があります")
    sure = True
    if shipper and suppliers:
        shipper, sure = match_supplier(shipper, suppliers)
    return {"decl": decl, "permit": permit, "shipper": shipper, "shipper_sure": sure}


def load_suppliers(root: Path) -> list[str]:
    """仕入先名の一覧: 台帳(全期)の仕入先名 + 仕入先名.txt(1行1社)。"""
    names = set()
    f = root / "仕入先名.txt"
    if f.exists():
        names |= {l.strip() for l in f.read_text(encoding="utf8").splitlines() if l.strip()}
    for lg in root.glob("第*期/輸入明細一覧_*.xlsx"):
        wb = openpyxl.load_workbook(lg, read_only=True)
        for ws, col in ((wb[SHEET_SEND], 4), (wb[SHEET_NOSEND], 3)):
            for r in ws.iter_rows(min_row=4, min_col=col, max_col=col, values_only=True):
                if r[0]:
                    names.add(str(r[0]).strip())
    return sorted(names)


# ---------- 台帳 ----------
def new_ledger(path: Path, period: int):
    wb = openpyxl.load_workbook(TPL_LEDGER)
    for ws in wb:
        # 見本データを消し、見出しだけ残す(「例」行は送金ありのみ残す)
        first = 5 if ws.title == SHEET_SEND else 4
        for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
            for c in row:
                c.value = None
    wb[SHEET_SEND]["B1"] = period_label(period)
    wb[SHEET_NOSEND]["A1"] = period_label(period)
    wb.save(path)


def next_no(ws, col: int, prefix: str, period: int) -> str:
    pat = re.compile(rf"^{re.escape(prefix)}{period}-(\d+)$")
    nums = [int(m.group(1)) for r in range(5 if ws.title == SHEET_SEND else 4, ws.max_row + 1)
            if (v := ws.cell(r, col).value) and (m := pat.match(str(v)))]
    return f"{prefix}{period}-{max(nums, default=0) + 1:03d}"


def next_row(ws, key_col: int) -> int:
    r = ws.max_row
    while r >= 4 and ws.cell(r, key_col).value is None and ws.cell(r, key_col + 1).value is None:
        r -= 1
    return max(r + 1, 5 if ws.title == SHEET_SEND else 4)


def add_send(ws, no, office, info):
    r = next_row(ws, 3)
    vals = {2: office, 3: no, 13: info["decl"], 14: info["shipper"], 4: info["shipper"],
            15: info["permit"]}
    for c, v in vals.items():
        ws.cell(r, c).value = v
    return r


def add_nosend(ws, no, office, info):
    r = next_row(ws, 2)
    vals = {1: office, 2: no, 3: info["shipper"], 9: info["decl"], 10: info["shipper"],
            11: info["permit"]}
    for c, v in vals.items():
        ws.cell(r, c).value = v
    return r


# ---------- 表紙 ----------
def make_cover(out: Path, no: str, info: dict, nosend: str | None):
    wb = openpyxl.load_workbook(TPL_COVER)
    if nosend is None:
        wb.remove(wb.worksheets[1])
        ws = wb.worksheets[0]
        ws["E16"] = None  # 見本の「乙仲 No.」を消す
        ymd = ("E10", "I10", "K10")
    else:
        wb.remove(wb.worksheets[0])
        ws = wb.worksheets[0]
        ymd = ("E10", "H10", "J10")
        if nosend in ("着払", "無償"):
            ws["H1" if nosend == "着払" else "J1"].fill = HILITE
    ws["E3"] = no
    ws["E5"] = info["decl"].replace(" ", "")
    ws["E7"] = info["shipper"]
    if info["permit"]:
        for cell, v in zip(ymd, (info["permit"].year, info["permit"].month, info["permit"].day)):
            ws[cell] = v
    wb.save(out)


# ---------- main ----------
def collect(paths):
    for p in map(Path, paths):
        if p.is_dir():
            yield from sorted(p.rglob("*.pdf"))
        else:
            yield p


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", nargs="+", help="輸入許可通知書の PDF またはそのフォルダ")
    ap.add_argument("--office", default="", help="事業所名/担当(例: 林六／東京)")
    ap.add_argument("--nosend", choices=["着払", "無償", "乙仲"], help="海外送金なしの場合の種別")
    ap.add_argument("--period", type=int, help="期(省略時は今日の日付から判定)")
    ap.add_argument("--out", type=Path, help="出力先(省略時はデスクトップ/輸入事後調査)")
    a = ap.parse_args()

    root = a.out or desktop() / "輸入事後調査"
    ok = 0
    suppliers = load_suppliers(root)
    for pdf in collect(a.pdf):
        try:
            info = read_permit(pdf, suppliers)
        except Exception as e:
            print(f"[NG] {pdf.name}: {e}", file=sys.stderr)
            continue
        period = a.period or fiscal_period(datetime.now())  # 期は処理日基準(許可日が前期のこともある)
        folder = root / f"第{period}期"
        (folder / "表紙").mkdir(parents=True, exist_ok=True)
        ledger = folder / f"輸入明細一覧_{period}期.xlsx"
        if not ledger.exists():
            new_ledger(ledger, period)
        wb = openpyxl.load_workbook(ledger)
        if a.nosend:
            ws, prefix = wb[SHEET_NOSEND], "経"
            no = next_no(ws, 2, prefix, period)
            row = add_nosend(ws, no, a.office, info)
        else:
            ws = wb[SHEET_SEND]
            no = next_no(ws, 3, "", period)
            row = add_send(ws, no, a.office, info)
        # 同じ申告番号の重複登録を防ぐ
        col = 9 if a.nosend else 13
        dup = [r for r in range(4, ws.max_row + 1)
               if r != row and str(ws.cell(r, col).value or "").replace(" ", "") == info["decl"].replace(" ", "")]
        if dup:
            print(f"[SKIP] {pdf.name}: 申告番号 {info['decl']} は台帳の {dup[0]} 行目に登録済み", file=sys.stderr)
            continue
        if not info["shipper_sure"]:
            for c in ((row, 3), (row, 10)) if a.nosend else ((row, 4), (row, 14)):
                ws.cell(*c).fill = HILITE
            print(f"      仕出人『{info['shipper']}』は照合できませんでした(黄色の欄を要確認)")
        wb.save(ledger)
        make_cover(folder / "表紙" / f"表紙_{no}.xlsx", no, info, a.nosend)
        ok += 1
        print(f"[OK] {pdf.name} -> {no}  申告番号 {info['decl']}  許可日 {info['permit']:%Y/%m/%d}" if info["permit"]
              else f"[OK] {pdf.name} -> {no}  申告番号 {info['decl']}  (許可日は手入力)")
    print(f"{ok} 件を {root} に保存しました")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
